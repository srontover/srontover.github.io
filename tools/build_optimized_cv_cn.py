from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("Xuehao_Xu_Research_CV_CN_v4_editable.docx")
NAVY = "12314D"
BLUE = "1E6E96"
INK = "1D2730"
MUTED = "5E6C77"
RULE = "C8D4DE"


def set_font(run, size, color=INK, bold=False, italic=False):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def set_para(paragraph, before=0, after=0, line=1.0, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep


def add_bottom_border(paragraph, color=RULE):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.first_child_found_in("w:pBdr")
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def add_section(doc, title):
    p = doc.add_paragraph()
    set_para(p, before=5, after=3, line=1.0, keep=True)
    set_font(p.add_run(title), 9.2, BLUE, bold=True)
    add_bottom_border(p)


def add_body(doc, text, before=0, after=2.5, size=9.15, color=INK, italic=False, keep=False):
    p = doc.add_paragraph()
    set_para(p, before=before, after=after, line=1.0, keep=keep)
    set_font(p.add_run(text), size, color, italic=italic)
    return p


def add_project_title(doc, title, date, subtitle):
    p = doc.add_paragraph()
    set_para(p, before=2, after=0.4, line=1.0, keep=True)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.12), WD_TAB_ALIGNMENT.RIGHT)
    set_font(p.add_run(title), 10.05, NAVY, bold=True)
    set_font(p.add_run(f"\t{date}"), 9.05, MUTED, italic=True)
    subtitle_para = doc.add_paragraph()
    set_para(subtitle_para, before=0, after=0.6, line=1.0, keep=True)
    set_font(subtitle_para.add_run(subtitle), 8.9, MUTED, italic=True)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_para(p, before=0, after=0.5, line=1.0)
    p.paragraph_format.left_indent = Inches(0.19)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    set_font(p.add_run(text), 8.55)


def add_skill(doc, label, value):
    p = doc.add_paragraph()
    set_para(p, before=0, after=0.8, line=1.0)
    set_font(p.add_run(f"{label}："), 8.6, NAVY, bold=True)
    set_font(p.add_run(value), 8.6)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.34)
    section.bottom_margin = Inches(0.12)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.08)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(8.7)
    normal.font.color.rgb = RGBColor.from_string(INK)

    title = doc.add_paragraph()
    set_para(title, after=0.5, line=1.0)
    set_font(title.add_run("徐学颢"), 21, NAVY, bold=True)
    set_font(title.add_run("  |  科研简历"), 10.2, MUTED)

    contact = doc.add_paragraph()
    set_para(contact, after=3, line=1.0)
    set_font(contact.add_run("邮箱: xuxue6800@gmail.com | GitHub: github.com/srontover | 个人主页: srontover.github.io"), 8.85, MUTED)
    add_bottom_border(contact, BLUE)

    add_section(doc, "研究兴趣")
    add_body(doc, "人形机器人学习；强化学习与学习控制；运动先验与模仿学习；统一技能策略；全身控制；接触丰富运动技能；仿真到真实迁移", after=2.5, size=9.05)

    add_section(doc, "教育经历")
    p = doc.add_paragraph()
    set_para(p, after=0.5, line=1.0, keep=True)
    set_font(p.add_run("南京邮电大学（NJUPT）"), 10.0, NAVY, bold=True)
    add_body(doc, "人工智能专业 工学学士候选人 | 预计毕业：2028 年", after=0.6, size=9.05, color=MUTED)
    add_body(doc, "GPA：4.08/5.0 | 人工智能专业排名：3/133", after=2, size=9.05)

    add_section(doc, "科研经历")
    projects = [
        (
            "人形机器人学习与运动先验研究",
            "2026.09-至今",
            "香港中文大学 | 科研实习生（远程） | 指导：颜子轲博士（清华大学 AIR）、李钟毓助理教授（香港中文大学机械与自动化工程学系）",
            [
                "围绕人形机器人学习、全身控制、运动模仿 / 运动先验与任务适应开展科研训练，包括文献调研、论文精读、Slide 汇报、研究问题提炼、方法讨论与实验规划。",
                "系统调研 RoboNaldo、Humanoid Goalkeeper、PAiD、PGMT、SMP、ADD、ConsMimic、CMDP 等工作，重点分析 task objective 与 motion imitation / motion prior 的冲突、条件化放松与多目标优化思路。",
                "当前关注 unified policy、非连续运动数据导致的 skill transition / OOD，以及不严重依赖 joint-level tracking error 的 human-like motion guidance；保持组内周级和 PI 双周级讨论，并计划复现 SMP 及足球相关基线。",
            ],
        ),
        (
            "Booster T1 人形机器人强化学习与动态技能训练",
            "2026-至今",
            "独立项目 | 步态、踢球与 Conditional AMP 探索",
            [
                "基于 23 DoF Booster T1，在 MuJoCo、MJLab 与 Warp 仿真工具链中实现并修改人形机器人强化学习训练环境，使用 PPO 训练 locomotion 与 kicking skill。",
                "配置 8192 个并行环境，完成观测与动作空间设计、球体重置、接触检测、奖励塑形、终止条件和课程学习等任务组件的实现与评估。",
                "已建立“接近足球-调整站位-向前踢球-恢复稳定姿态”的训练与评估流程，使用接触质量、球速、目标方向误差、终止原因与训练稳定性分析策略行为；高精度射门误差优化仍在进行。",
                "开展 Conditional AMP / motion-prior 探索，研究 expert motion 与 policy condition 的对齐、动作泛化及训练稳定性。",
            ],
        ),
        (
            "基于三维感知与几何约束的机器人手眼标定",
            "2026",
            "南京邮电大学 | 科研项目 | 论文工作推进中",
            [
                "面向工业机器人与三维视觉系统开展手眼标定研究，使用 ABB IRB-120 / IRC5 和 Keyence LJ-X8000A / LJ-X8060 完成机器人运动、IO 触发、轮廓采集与系统联调。",
                "设计并验证包含正交几何约束预标定、动态点云筛选、精细优化和独立测试的多阶段标定流程。",
                "完成多姿态、多边缘、多点数据采集，并参与处理机器人零点、SMB 电池、触发同步等实际系统问题；相关发明专利申请已公开：CN121821412A，论文工作仍在推进中。",
            ],
        ),
    ]
    for title_text, date, subtitle, bullets in projects:
        add_project_title(doc, title_text, date, subtitle)
        for bullet in bullets:
            add_bullet(doc, bullet)

    add_section(doc, "竞赛工程经历")
    add_project_title(doc, "简易自行瞄准装置", "2024", "全国大学生电子设计竞赛 | 视觉感知与目标跟踪负责人 | 3 人团队 | 江苏省三等奖")
    for bullet in [
        "基于 Raspberry Pi 与工业摄像头搭建实时视觉处理系统，负责目标检测、靶心定位与持续跟踪模块。",
        "使用 OpenCV 实现颜色阈值分割、轮廓提取与圆形检测，以约 30 Hz 向下游控制模块输出靶心位置信息；调试阈值与目标筛选逻辑以提升移动小车和云台运动下的跟踪稳定性，并与云台及执行机构控制模块联调，支持闭环目标跟踪和自动瞄准。",
    ]:
        add_bullet(doc, bullet)

    add_section(doc, "技术能力")
    skills = [
        ("编程与工程", "Python、C++、Git/GitHub、Linux"),
        ("机器学习与机器人学习", "PyTorch、PPO、奖励设计、课程学习、人形机器人运动控制、运动模仿、运动先验、Conditional AMP"),
        ("仿真、机器人与视觉系统", "MuJoCo、MJLab、Warp、ABB RAPID、机器人标定、三维感知、点云处理、OpenCV、目标跟踪、相机-控制接口联调"),
    ]
    for label, value in skills:
        add_skill(doc, label, value)

    add_section(doc, "竞赛与荣誉")
    for award in [
        "RoboCup 中国公开赛 3D 仿真组第三名（2025）。",
        "江苏省机器人竞赛 3D 仿真组第三名（2025）。",
    ]:
        add_bullet(doc, award)

    doc.core_properties.title = "徐学颢 - 科研简历"
    doc.core_properties.author = "徐学颢"
    doc.core_properties.subject = "科研简历"
    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
