from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("Xuehao_Xu_Research_CV_EN_v4_editable.docx")
NAVY = "12314D"
BLUE = "1E6E96"
INK = "1D2730"
MUTED = "5E6C77"
RULE = "C8D4DE"


def set_font(run, size, color=INK, bold=False, italic=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def set_para(paragraph, before=0, after=0, line=1.0, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep


def add_bottom_border(paragraph, color=RULE):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.first_child_found_in("w:pBdr")
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def add_section(doc, title):
    p = doc.add_paragraph()
    set_para(p, before=5, after=3, line=1.0, keep=True)
    set_font(p.add_run(title), 9.2, BLUE, bold=True)
    add_bottom_border(p)


def add_body(doc, text, before=0, after=2.5, size=9.1, color=INK, italic=False, keep=False):
    p = doc.add_paragraph()
    set_para(p, before=before, after=after, line=1.0, keep=keep)
    set_font(p.add_run(text), size, color, italic=italic)
    return p


def add_project_title(doc, title, date, subtitle):
    p = doc.add_paragraph()
    set_para(p, before=2, after=0.4, line=1.0, keep=True)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.12), WD_TAB_ALIGNMENT.RIGHT)
    set_font(p.add_run(title), 10.05, NAVY, bold=True)
    set_font(p.add_run(f"\t{date}"), 9.05, MUTED, italic=True)
    subtitle_para = doc.add_paragraph()
    set_para(subtitle_para, before=0, after=0.6, line=1.0, keep=True)
    set_font(subtitle_para.add_run(subtitle), 8.8, MUTED, italic=True)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_para(p, before=0, after=0.5, line=1.0)
    p.paragraph_format.left_indent = Inches(0.19)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    set_font(p.add_run(text), 8.75)


def add_skill(doc, label, value):
    p = doc.add_paragraph()
    set_para(p, before=0, after=0.8, line=1.0)
    set_font(p.add_run(f"{label}: "), 8.8, NAVY, bold=True)
    set_font(p.add_run(value), 8.8)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.38)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(8.9)
    normal.font.color.rgb = RGBColor.from_string(INK)

    title = doc.add_paragraph()
    set_para(title, after=0.5, line=1.0)
    set_font(title.add_run("XUEHAO XU"), 21, NAVY, bold=True)
    set_font(title.add_run("  |  Research CV"), 10.2, MUTED)

    contact = doc.add_paragraph()
    set_para(contact, after=3, line=1.0)
    set_font(contact.add_run("Email: xuxue6800@gmail.com | GitHub: github.com/srontover | Homepage: srontover.github.io"), 8.85, MUTED)
    add_bottom_border(contact, BLUE)

    add_section(doc, "RESEARCH INTERESTS")
    add_body(doc, "Humanoid Robot Learning; Reinforcement Learning and Learning-Based Control; Motion Priors and Imitation Learning; Unified Skill Policies; Whole-Body Control; Contact-Rich Physical Skills; Sim-to-Real Transfer", after=2.5, size=9.0)

    add_section(doc, "EDUCATION")
    p = doc.add_paragraph()
    set_para(p, after=0.5, line=1.0, keep=True)
    set_font(p.add_run("Nanjing University of Posts and Telecommunications (NJUPT)"), 10.0, NAVY, bold=True)
    add_body(doc, "B.Eng. Candidate in Artificial Intelligence | Expected Graduation: 2028", after=0.6, size=9.05, color=MUTED)
    add_body(doc, "GPA: 4.08/5.0 | Ranked 3/133 in Artificial Intelligence", after=2, size=9.05)

    add_section(doc, "RESEARCH EXPERIENCE")
    projects = [
        (
            "Humanoid Robot Learning and Motion Prior Research",
            "Sep 2026-Present",
            "The Chinese University of Hong Kong | Research Intern (Remote) | Supervised by Dr. Zike Yan (Tsinghua AIR) and Prof. Zhongyu Li (CUHK MAE)",
            [
                "Participating in research training on humanoid robot learning, whole-body control, motion imitation and motion priors, and task adaptation through literature review, paper presentations, research-question formulation, and experiment planning.",
                "Reviewed RoboNaldo, Humanoid Goalkeeper, PAiD, PGMT, SMP, ADD, ConsMimic, and CMDP, analyzing conflicts between task objectives and motion imitation or priors, conditional relaxation, and multi-objective formulations.",
                "Exploring unified policies, skill-transition and OOD issues from discontinuous motion data, and human-like motion guidance beyond joint-level tracking errors; planning reproduction of SMP and soccer-related baselines.",
            ],
        ),
        (
            "Reinforcement Learning for Dynamic Humanoid Skills",
            "2026-Present",
            "Independent Project | Booster T1 Locomotion, Kicking, and Conditional AMP Exploration",
            [
                "Implemented and modified reinforcement-learning environments for a 23-DoF Booster T1 humanoid using MuJoCo, MJLab, and Warp; trained locomotion and kicking skills with PPO.",
                "Configured 8,192 parallel environments and implemented observation and action spaces, ball-reset logic, contact detection, reward shaping, termination conditions, and curriculum learning.",
                "Established a training and evaluation workflow for approach, stance adjustment, forward kicking, and recovery; analyzed contact quality, ball velocity, target-direction error, termination causes, and training stability. High-precision kicking accuracy remains under optimization.",
                "Investigating Conditional AMP and motion-prior approaches for expert-motion and policy-condition alignment, skill generalization, and training stability.",
            ],
        ),
        (
            "Robot Hand-Eye Calibration with 3D Sensing and Geometric Constraints",
            "2026",
            "Nanjing University of Posts and Telecommunications | Research Project | Manuscript in Preparation",
            [
                "Conducted hand-eye calibration research for industrial robot and 3D vision systems using an ABB IRB-120/IRC5 and Keyence LJ-X8000A/LJ-X8060 sensors, including robot motion, I/O triggering, profile acquisition, and system integration.",
                "Designed and evaluated a multi-stage calibration workflow with orthogonal geometric pre-calibration, dynamic point-cloud screening, refinement optimization, and independent testing.",
                "Collected multi-pose, multi-edge, and multi-point data and helped address practical system issues including robot zeroing, SMB battery maintenance, and trigger synchronization. Published patent application: CN121821412A; manuscript in preparation.",
            ],
        ),
    ]
    for title_text, date, subtitle, bullets in projects:
        add_project_title(doc, title_text, date, subtitle)
        for bullet in bullets:
            add_bullet(doc, bullet)

    add_section(doc, "COMPETITION ENGINEERING EXPERIENCE")
    add_project_title(doc, "Autonomous Aiming Device", "2025", "National Undergraduate Electronic Design Contest | Vision Perception and Target Tracking | Three-Member Team | Jiangsu Provincial Third Prize")
    for bullet in [
        "Built a real-time visual-processing pipeline with a Raspberry Pi and industrial camera; responsible for target detection, target-center localization, and continuous tracking.",
        "Implemented color-threshold segmentation, contour extraction, and circle detection in OpenCV with approximately 30 Hz feedback; tuned thresholds and target-selection logic for vehicle and gimbal motion, then integrated the pipeline with gimbal and actuator-control modules for closed-loop aiming.",
    ]:
        add_bullet(doc, bullet)

    add_section(doc, "TECHNICAL SKILLS")
    for label, value in [
        ("Programming and Engineering", "Python, C++, Git/GitHub, Linux"),
        ("ML and Robot Learning", "PyTorch, PPO, reward design, curriculum learning, humanoid motion control, motion imitation, motion priors, Conditional AMP"),
        ("Simulation, Robotics, and Vision", "MuJoCo, MJLab, Warp, ABB RAPID, robot calibration, 3D sensing, point-cloud processing, OpenCV, target tracking, camera-control integration"),
    ]:
        add_skill(doc, label, value)

    add_section(doc, "HONORS")
    for award in [
        "Third Place, RoboCup China Open 3D Simulation League, 2025.",
        "Third Place, Jiangsu Provincial Robot Competition, 3D Simulation Group, 2025.",
    ]:
        add_bullet(doc, award)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(footer, before=0, after=0, line=1.0)
    set_font(footer.add_run("Xuehao Xu | Research CV"), 7.5, MUTED)

    doc.core_properties.title = "Xuehao Xu - Research CV"
    doc.core_properties.author = "Xuehao Xu"
    doc.core_properties.subject = "Research CV"
    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
